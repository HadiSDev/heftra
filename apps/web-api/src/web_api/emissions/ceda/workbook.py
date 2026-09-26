"""Read an Open CEDA release (the 2025 layout) into plain values."""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from ..countries import alpha2
from ..factors import REST_OF_WORLD
from .sheets import Row, decimal, header_index, labelled_value, sheet_rows, text
from .types import CedaWorkbook, WorkbookError, WorkbookSector

COVER = "Cover"
FACTORS = "GHG_t_Raw"
REGIONS = "Regional Average EFs"
PURCHASER_RATIOS = "Purchaser - producer conversion"
COUNTRY_REGIONS = "Country to region mapping"
METADATA = "Metadata"

REST_OF_WORLD_CODE = "ROW"

CURRENCIES = {"US Dollar": "USD", "USD": "USD"}
PRICE_TYPES = {"Producer price": "producer", "Purchaser price": "purchaser"}
REGION_SPELLINGS = {
    "Australian and New Zealand": "Australia and New Zealand",
    "Carribean": "Caribbean",
    "South Eastern Asia": "South-Eastern Asia",
}


def read_workbook(path: Path) -> CedaWorkbook:
    """Everything an import needs from the workbook at `path`."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return _read(
            cover=sheet_rows(workbook, COVER),
            factors=sheet_rows(workbook, FACTORS),
            regions=sheet_rows(workbook, REGIONS),
            ratios=sheet_rows(workbook, PURCHASER_RATIOS),
            country_regions=sheet_rows(workbook, COUNTRY_REGIONS),
            metadata=sheet_rows(workbook, METADATA),
        )
    finally:
        workbook.close()


def _read(*, cover: list[Row], factors: list[Row], regions: list[Row], ratios: list[Row],
          country_regions: list[Row], metadata: list[Row]) -> CedaWorkbook:
    currency = _currency(factors, FACTORS)
    price_year = _year(factors, FACTORS)
    if (_currency(regions, REGIONS), _year(regions, REGIONS)) != (currency, price_year):
        raise WorkbookError(f"{REGIONS!r} is not in the currency and year of {FACTORS!r}")

    country_factors, region_factors, skipped = _country_factors(factors)
    region_factors.update(_region_factors(regions))
    return CedaWorkbook(
        version=labelled_value(cover, "Version", COVER),
        currency=currency,
        price_year=price_year,
        price_type=_price_type(factors),
        sectors=_sectors(metadata),
        country_factors=country_factors,
        region_factors=region_factors,
        purchaser_ratios=_purchaser_ratios(ratios),
        country_regions=_country_regions(country_regions),
        skipped_countries=skipped,
    )


def _currency(rows: list[Row], sheet: str) -> str:
    label = labelled_value(rows, "Currency", sheet)
    if label not in CURRENCIES:
        raise WorkbookError(f"sheet {sheet!r} is in an unsupported currency: {label!r}")
    return CURRENCIES[label]


def _year(rows: list[Row], sheet: str) -> int:
    value = labelled_value(rows, "Year", sheet)
    if not value.isdigit():
        raise WorkbookError(f"sheet {sheet!r} has no year: {value!r}")
    return int(value)


def _price_type(rows: list[Row]) -> str:
    label = labelled_value(rows, "Price type", FACTORS)
    if label not in PRICE_TYPES:
        raise WorkbookError(f"sheet {FACTORS!r} has an unsupported price type: {label!r}")
    return PRICE_TYPES[label]


def _codes(row: Row, first_column: int) -> list[str]:
    codes = [text(cell) for cell in row[first_column:]]
    named = [code for code in codes if code]
    duplicates = sorted({code for code in named if named.count(code) > 1})
    if duplicates:
        raise WorkbookError(f"sectors listed twice: {duplicates}")
    return codes


def _country_factors(
    rows: list[Row],
) -> tuple[dict[tuple[str, str], Decimal], dict[tuple[str, str], Decimal], list[str]]:
    header = header_index(rows, "Country Code", FACTORS)
    codes = _codes(rows[header], 3)
    by_country: dict[tuple[str, str], Decimal] = {}
    rest_of_world: dict[tuple[str, str], Decimal] = {}
    skipped: list[str] = []
    seen: set[str] = set()
    for row in rows[header + 1:]:
        country = text(row[0])
        if not country:
            break
        if country in seen:
            raise WorkbookError(f"{FACTORS!r} lists {country} twice")
        seen.add(country)
        values = _values(row, codes, 3, f"{FACTORS!r} row {country}")
        if country == REST_OF_WORLD_CODE:
            rest_of_world.update({(code, REST_OF_WORLD): value for code, value in values.items()})
            continue
        code2 = alpha2(country)
        if code2 is None:
            skipped.append(country)
            continue
        by_country.update({(code, code2): value for code, value in values.items()})
    return by_country, rest_of_world, skipped


def _region_factors(rows: list[Row]) -> dict[tuple[str, str], Decimal]:
    header = header_index(rows, "Region", REGIONS)
    codes = _codes(rows[header], 2)
    factors: dict[tuple[str, str], Decimal] = {}
    for row in rows[header + 1:]:
        region = _region(text(row[0]))
        if not region:
            break
        values = _values(row, codes, 2, f"{REGIONS!r} row {region}")
        factors.update({(code, region): value for code, value in values.items()})
    return factors


def _values(row: Row, codes: list[str], first_column: int, where: str) -> dict[str, Decimal]:
    return {
        code: decimal(value, f"{where}, sector {code}")
        for code, value in zip(codes, row[first_column:])
        if code
    }


def _purchaser_ratios(rows: list[Row]) -> dict[str, Decimal]:
    header = header_index(rows, "Sector Code", PURCHASER_RATIOS)
    values = rows[header_index(rows, "Purchaser - Producer conversion", PURCHASER_RATIOS,
                               after=header)]
    return _values(values, _codes(rows[header], 1), 1, f"{PURCHASER_RATIOS!r} ratio")


def _country_regions(rows: list[Row]) -> dict[str, str]:
    header = header_index(rows, "Country Code", COUNTRY_REGIONS)
    regions: dict[str, str] = {}
    for row in rows[header + 1:]:
        country = text(row[0])
        if not country:
            break
        code2 = alpha2(country)
        if code2 is not None:
            regions[code2] = _region(text(row[1]))
    return regions


def _region(name: str) -> str:
    return REGION_SPELLINGS.get(name, name)


def _sectors(rows: list[Row]) -> list[WorkbookSector]:
    header = header_index(rows, "Sector Code", METADATA)
    sectors: list[WorkbookSector] = []
    for row in rows[header + 1:]:
        code = text(row[0])
        if not code:
            break
        sectors.append(WorkbookSector(code, text(row[1]), text(row[2]) or None))
    return sectors
